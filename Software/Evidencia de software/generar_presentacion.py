from io import BytesIO
from pathlib import Path
import sys

SCRIPT_DIR = Path(__file__).resolve().parent
VENV_PACKAGES = (
    SCRIPT_DIR.parent / ".venv312" / "lib"
    / f"python{sys.version_info.major}.{sys.version_info.minor}" / "site-packages"
)
if VENV_PACKAGES.is_dir() and str(VENV_PACKAGES) not in sys.path:
    sys.path.insert(0, str(VENV_PACKAGES))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
CAPTURES = ROOT / "Capturas"
LOGO = ROOT.parent.parent / "Logo_DuocUC.svg.webp"
OUTPUT = ROOT / "Presentacion_Demo_Flota_Bomberos.pptx"

NAVY = RGBColor(20, 43, 61)
RED = RGBColor(198, 47, 51)
GOLD = RGBColor(225, 166, 45)
PAPER = RGBColor(246, 248, 250)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(35, 48, 60)
MUTED = RGBColor(91, 105, 117)
GREEN = RGBColor(34, 126, 91)
LIGHT_GREEN = RGBColor(229, 244, 237)
LINE = RGBColor(222, 228, 233)

SLIDE_W = 13.333
SLIDE_H = 7.5
IMAGE_X = 0.55
IMAGE_Y = 1.38
IMAGE_W = 8.65
IMAGE_H = IMAGE_W * 9 / 16
PANEL_X = 9.45
PANEL_Y = IMAGE_Y
PANEL_W = 3.33
PANEL_H = IMAGE_H


def logo_bytes():
    with Image.open(LOGO) as source:
        image = source.convert("RGBA")
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()


LOGO_PNG = logo_bytes()


def add_text(slide, text, x, y, w, h, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Arial"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_logo(slide, x, y, width):
    slide.shapes.add_picture(BytesIO(LOGO_PNG), Inches(x), Inches(y), width=Inches(width))


def add_background(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(0.11)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED


def add_header(slide, number, title, subtitle=""):
    add_background(slide)
    add_text(slide, f"EVIDENCIA FUNCIONAL  ·  {number:02d}", 0.58, 0.24, 4.6, 0.22,
             size=9, color=RED, bold=True)
    add_text(slide, title, 0.58, 0.48, 10.2, 0.46, size=22, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.6, 0.94, 10.1, 0.25, size=10, color=MUTED)
    add_logo(slide, 11.48, 0.28, 1.2)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.22), Inches(12.2), Inches(0.012)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = LINE
    add_footer(slide, number)


def add_footer(slide, number):
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(7.03), Inches(12.2), Inches(0.012)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = LINE
    add_text(slide, "DEMO  ·  Base capstone  ·  Registros ficticios", 0.62, 7.08, 8.4, 0.18,
             size=8, color=MUTED)
    add_text(slide, f"{number:02d}", 12.15, 7.07, 0.56, 0.18,
             size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_screenshot(slide, filename):
    image = CAPTURES / filename
    if not image.is_file():
        raise FileNotFoundError(f"No existe la captura requerida: {image}")
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(IMAGE_X - 0.035), Inches(IMAGE_Y - 0.035),
        Inches(IMAGE_W + 0.07), Inches(IMAGE_H + 0.07),
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE
    slide.shapes.add_picture(
        str(image), Inches(IMAGE_X), Inches(IMAGE_Y),
        width=Inches(IMAGE_W), height=Inches(IMAGE_H),
    )


def add_panel(slide, heading, points, accent=RED):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(PANEL_X), Inches(PANEL_Y), Inches(PANEL_W), Inches(PANEL_H),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(PANEL_X + 0.2), Inches(PANEL_Y + 0.27),
        Inches(0.08), Inches(0.3),
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    add_text(slide, heading, PANEL_X + 0.38, PANEL_Y + 0.21, PANEL_W - 0.58, 0.42,
             size=15, color=NAVY, bold=True)
    top = PANEL_Y + 0.86
    for title, body in points:
        add_text(slide, title, PANEL_X + 0.28, top, PANEL_W - 0.52, 0.28,
                 size=11, color=RED, bold=True)
        top += 0.29
        add_text(slide, body, PANEL_X + 0.28, top, PANEL_W - 0.52, 0.61,
                 size=10.5, color=INK)
        top += 0.73


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_logo(slide, 0.78, 0.62, 2.65)
    add_text(slide, "PRESENTACIÓN DE EVIDENCIAS", 0.82, 1.63, 4.25, 0.3,
             size=10, color=GOLD, bold=True)
    add_text(slide, "Gestión de flota\ny mantenciones", 0.8, 2.0, 4.35, 1.35,
             size=28, color=WHITE, bold=True)
    add_text(slide, "Bomberos de Talcahuano  ·  Odoo 18", 0.83, 3.5, 4.15, 0.45,
             size=16, color=WHITE)
    chip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(4.22), Inches(2.0), Inches(0.42)
    )
    chip.fill.solid()
    chip.fill.fore_color.rgb = RED
    chip.line.fill.background()
    add_text(slide, "DATOS DE DEMOSTRACIÓN", 0.94, 4.27, 1.78, 0.27,
             size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    cover_image = CAPTURES / "06_flota.png"
    slide.shapes.add_picture(
        str(cover_image), Inches(5.35), Inches(1.32),
        width=Inches(7.35), height=Inches(7.35 * 9 / 16),
    )
    add_text(slide, "Pruebas funcionales ejecutadas en la base activa capstone",
             0.84, 5.1, 4.15, 0.75, size=13, color=RGBColor(225, 232, 238))
    add_text(slide, "14 de octubre de 2026", 0.84, 6.58, 4.2, 0.24,
             size=9, color=RGBColor(194, 208, 219))
    add_text(slide, "Odoo 18  ·  3 addons propios  ·  9 pruebas aprobadas",
             5.38, 5.62, 7.3, 0.32, size=11, color=WHITE, bold=True)
    return slide


def add_image_slide(prs, number, title, subtitle, capture, heading, points, accent=RED):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, number, title, subtitle)
    add_screenshot(slide, capture)
    add_panel(slide, heading, points, accent)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "Gestión de flota y mantenciones - Bomberos de Talcahuano"
    prs.core_properties.subject = "Evidencias de software y pruebas funcionales en Odoo 18"
    prs.core_properties.author = "Proyecto Capstone - Ingeniería en Informática"
    add_cover(prs)

    add_image_slide(
        prs, 2, "Ingreso al sistema", "La base seleccionada para esta demostración es capstone.",
        "01_ingreso.png", "Cómo se ingresa",
        [
            ("01  BASE", "Se selecciona la base capstone."),
            ("02  USUARIO", "El acceso depende del perfil asignado por el administrador."),
            ("03  MÓDULO", "Desde el menú se abre Bomberos para entrar a Flota o Mantenciones."),
            ("CUIDADO", "La presentación no incluye contraseñas ni credenciales reales."),
        ],
    )
    add_image_slide(
        prs, 3, "Roles de la aplicación", "Grupos propios de Bomberos, no aplicaciones empresariales reutilizadas.",
        "03_roles.png", "Permisos por perfil",
        [
            ("ADMINISTRADOR", "Consulta toda la información de Bomberos."),
            ("JEFE DE FLOTA", "Ve todas las compañías y puede aprobar/cerrar órdenes."),
            ("TÉCNICO", "Opera mantenciones y kilometrajes de sus compañías asignadas."),
            ("BOMBERO CONSULTOR", "Consulta vehículos, órdenes y avisos de sus compañías."),
        ],
    )
    add_image_slide(
        prs, 4, "Usuarios y compañías asignadas", "Los perfiles de ejemplo permiten mostrar la separación de acceso.",
        "12_perfil_consultor_compania.png", "Ejemplo de asignación",
        [
            ("4 PERFILES DEMO", "Administrador, Jefatura, Técnico y Bombero consultor."),
            ("CONSULTOR", "Este perfil de ejemplo está asignado solo a la Compañía Talcahuano 02."),
            ("TÉCNICO", "El perfil de ejemplo está asignado a la Compañía Talcahuano 01."),
            ("IMPORTANTE", "Las cuentas DEMO no tienen contraseña configurada; deben habilitarse antes de iniciar sesión."),
        ],
    )
    add_image_slide(
        prs, 5, "Flota y kilometraje", "Cinco vehículos ficticios repartidos entre dos compañías DEMO.",
        "06_flota.png", "Qué muestra la ficha",
        [
            ("IDENTIFICACIÓN", "Código, unidad, patente DEMO, compañía y tipo de vehículo."),
            ("KILOMETRAJE", "La cifra actual proviene del historial de lecturas."),
            ("ESTADO", "Una unidad con una orden en ejecución aparece En mantención."),
            ("DATOS", "Los códigos y lecturas son de muestra; no describen vehículos reales."),
        ],
    )
    add_image_slide(
        prs, 6, "Órdenes y flujo de mantención", "Las órdenes DEMO hacen visible cada etapa del proceso.",
        "07_ordenes.png", "Flujo y responsables",
        [
            ("BORRADOR → PROGRAMADA", "Se registra la necesidad y se agenda la intervención."),
            ("EN EJECUCIÓN", "El Técnico registra diagnóstico, trabajo, kilometraje y recursos."),
            ("EN REVISIÓN", "El Técnico solicita aprobación cuando termina el trabajo."),
            ("CERRADA", "Solo la Jefatura de flota puede cerrar la orden."),
        ],
    )
    add_image_slide(
        prs, 7, "Materiales, mano de obra y costos", "Ejemplo de orden reactiva DEMO-MT-00002.",
        "13_materiales_costos.png", "Cálculo demostrado",
        [
            ("MATERIALES", "$25.000 CLP en repuesto ficticio."),
            ("MANO DE OBRA", "$37.500 CLP (2,5 horas x $15.000)."),
            ("TOTAL", "$62.500 CLP, calculado por el sistema."),
            ("ALCANCE", "Se registran costos por orden; no se controla inventario de bodega."),
        ],
        accent=GOLD,
    )
    add_image_slide(
        prs, 8, "Cierre con trazabilidad", "La ficha conserva quién revisó y el kilometraje final.",
        "11_orden_cerrada.png", "Resultado de DEMO-MT-00005",
        [
            ("ESTADO", "Cerrada después de revisión."),
            ("RESPONSABLE", "El registro identifica al Jefe de flota que aprobó."),
            ("LECTURA FINAL", "12.600 km quedan en el historial del vehículo."),
            ("COSTO", "$22.000 CLP en este caso de demostración."),
        ],
        accent=GREEN,
    )
    add_image_slide(
        prs, 9, "Avisos preventivos y de revisión", "Avisos internos visibles en la aplicación.",
        "08_avisos.png", "Cuándo se avisa",
        [
            ("MANTENCIÓN PRÓXIMA", "Dentro de 30 días o a menos de 500 km del objetivo."),
            ("EJEMPLO", "DEMO-MT-00001 aparece como próxima por fecha y kilometraje."),
            ("REVISIÓN", "DEMO-MT-00003 genera un aviso mientras espera a Jefatura."),
            ("ACTUALIZACIÓN", "El proceso programado revisa los avisos diariamente."),
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, 10, "Pruebas automatizadas ejecutadas",
               "Suite dirigida a los tres addons propios, ejecutada en la base de trabajo capstone.")
    stats = [
        ("9", "PRUEBAS", NAVY),
        ("0", "FALLAS", GREEN),
        ("0", "ERRORES", GREEN),
    ]
    for index, (value, label, color) in enumerate(stats):
        x = 0.75 + index * 2.15
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.55), Inches(1.85), Inches(1.25)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = LINE
        add_text(slide, value, x + 0.12, 1.65, 1.6, 0.62, size=32, color=color,
                 bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.12, 2.3, 1.6, 0.25, size=9, color=MUTED,
                 bold=True, align=PP_ALIGN.CENTER)

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.35), Inches(1.55), Inches(5.25), Inches(4.8)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = LINE
    add_text(slide, "Cobertura comprobada", 7.72, 1.82, 4.45, 0.42,
             size=17, color=NAVY, bold=True)
    add_text(
        slide,
        "• Acceso por compañía y lectura global de Jefatura.\n"
        "• Kilometraje no decreciente y restricciones por rol.\n"
        "• Flujo de orden, revisión, cierre y protección de registros.\n"
        "• Costos, avisos próximos/vencidos y prevención de duplicados.",
        7.75, 2.48, 4.4, 2.45, size=13, color=INK,
    )
    add_text(slide, "ALCANCE", 7.75, 5.15, 1.0, 0.25, size=9, color=RED, bold=True)
    add_text(
        slide,
        "Validación automatizada de los addons. La prueba manual con usuarios Bomberos "
        "y dispositivos móviles/PWA sigue pendiente.",
        7.75, 5.43, 4.35, 0.62, size=10.5, color=MUTED,
    )
    add_text(slide, "Las capturas y registros con prefijo DEMO son datos ficticios de prueba.",
             0.82, 4.05, 5.85, 0.65, size=15, color=NAVY, bold=True)
    add_text(slide, "Las aplicaciones estándar Fleet, Maintenance y Stock no se instalaron.",
             0.82, 4.92, 5.85, 0.58, size=12, color=INK)
    add_text(slide, "El código fuente de Odoo permanece intacto; se utilizan addons propios.",
             0.82, 5.7, 5.85, 0.55, size=12, color=INK)

    prs.save(OUTPUT)
    print(f"Presentación creada: {OUTPUT}")
    print(f"Diapositivas: {len(prs.slides)}")


if __name__ == "__main__":
    build()
